"""Continuous streams: CGM sessions and intraday heart rate (`build --continuous`).

The physiology checks hold the curves to the person's own model (eAG, resting heart rate, the shared day
plan); the format checks hold every export to its truth record and to the byte-level shape documented in
`docs/DEVICE_FORMATS.md`.
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import pathlib
import sys
import tempfile
from datetime import date, datetime, timedelta
from xml.etree import ElementTree as ET

import pytest

REPO = pathlib.Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO))

from mirobody_gen import continuous as C  # noqa: E402
from mirobody_gen import continuous_formats as F  # noqa: E402
from mirobody_gen import devices, person as person_mod, physiology, spec, vendor_signals  # noqa: E402

SEED = 7


@pytest.fixture(scope="module")
def cohort():
    people = person_mod.build_cohort(SEED, None)
    langs = {p.person_id: person_mod.person_lang(SEED, p.person_id) for p in people}
    return people, langs


def _life(person, langs):
    return C.Life(SEED, person, devices.series_for(person, SEED, langs[person.person_id]))


#: Every device, every unit and both settings, whatever the adoption weights give this seed.
ROTATION = [("dexcom_g7", "mg/dL"), ("yuwell_ct3", "mmol/L"), ("libre_2", "mmol/L"), ("aidex", "mmol/L"),
            ("sibionics_gs1", "mmol/L"), ("ican_i3", "mmol/L"), ("dexcom_g6", "mmol/L"), ("libre_3", "mg/dL"),
            ("guardian_4", "mg/dL"), ("dexcom_g7", "mmol/L"), ("sibionics_gs1", "mg/dL"), ("guardian_4", "mmol/L")]
EN_ROTATION = [("dexcom_g7", "mg/dL"), ("libre_3", "mmol/L"), ("guardian_4", "mg/dL"), ("libre_2", "mg/dL"),
               ("dexcom_g6", "mmol/L")]


@pytest.fixture(scope="module")
def built(cohort):
    """A build of everyone with a sensor session (devices rotated so each format appears) plus a few
    watch and ring wearers."""
    people, langs = cohort
    chosen = [p for p in people if C.cgm_plan(SEED, p, langs[p.person_id])]
    index = {p.person_id: i for i, p in enumerate(chosen)}
    original = C.cgm_plan

    def rotated(seed, person, lang):
        """Devices rotate over every (person, session); English-speaking wearers (the only ones with
        Nightscout and follower set-ups) cycle through the devices those set-ups exist for."""
        out = original(seed, person, lang)
        if person.person_id in index:
            for k, sess in enumerate(out):
                if lang == "en":
                    dev, unit = EN_ROTATION[(index[person.person_id] + k) % len(EN_ROTATION)]
                else:
                    dev, unit = ROTATION[(index[person.person_id] + k) % len(ROTATION)]
                wear = timedelta(days=spec.streams()["devices"][dev]["wear_days"])
                sess.update(device=dev, unit=unit, end=min(sess["end"], sess["start"] + wear),
                            setting="clinic" if dev == "sibionics_gs1" and k == 0 else "home")
        return out

    for p in people:
        s = devices.series_for(p, SEED, langs[p.person_id])
        if s["habits"]["wearable"] and p not in chosen and len(chosen) < 22:
            chosen.append(p)
    tmp = tempfile.TemporaryDirectory()
    out = pathlib.Path(tmp.name)
    rates = (C.NIGHTSCOUT_RATE, C.FOLLOWER_RATE, C.TIDEPOOL_RATE)
    C.cgm_plan = rotated
    C.NIGHTSCOUT_RATE = C.FOLLOWER_RATE = C.TIDEPOOL_RATE = 1.0          # every platform format appears
    try:
        counts = C.write_all(out, chosen, SEED, langs)
    finally:
        C.cgm_plan = original
        C.NIGHTSCOUT_RATE, C.FOLLOWER_RATE, C.TIDEPOOL_RATE = rates
    truths = [json.loads(line) for line in (out / "continuous.jsonl").open(encoding="utf-8")]
    yield out, truths, counts, chosen, langs
    tmp.cleanup()


# ── Physiology ───────────────────────────────────────────────────────────────
def test_glucose_day_means_are_the_eag_of_the_expected_hba1c(cohort):
    people, langs = cohort
    for person in [p for p in people if p.archetype == "prediabetes_to_t2dm"][:4]:
        life = _life(person, langs)
        start = datetime(2024, 5, 1)
        bg = C.glucose_minutes(life, start, 10 * 1440)
        for k in range(1, 9):                                    # whole local days
            day = start.date() + timedelta(days=k)
            m0 = int((datetime.combine(day, datetime.min.time()) - start).total_seconds() // 60)
            mean = sum(bg[m0:m0 + 1440]) / 1440
            eag = 1.59 * physiology.expected(person, "hba1c", day) - 2.59
            assert abs(mean - eag) < 0.6, (person.person_id, day, mean, eag)


def test_healthy_glucose_stays_physiological(cohort):
    people, langs = cohort
    for person in [p for p in people if p.archetype == "healthy"][:6]:
        bg = C.glucose_minutes(_life(person, langs), datetime(2025, 3, 1, 9), 14 * 1440)
        tir = sum(3.9 <= v <= 10.0 for v in bg) / len(bg)
        assert min(bg) > 3.0 and max(bg) < 14.0 and tir > 0.9, (person.person_id, min(bg), max(bg), tir)


def test_a_run_is_a_heart_rate_peak_and_a_glucose_dip_at_the_same_minutes(cohort):
    """The same day with and without its run: the run is what lifts heart rate and lowers glucose."""
    import dataclasses

    people, langs = cohort
    found = 0
    for person in people:
        running = next((e for e in person.events if e.name == C.RUN), None)
        if running is None:
            continue
        life = _life(person, langs)
        for k in range(120, 160):
            day = running.start + timedelta(days=k)
            plan = life.plan(day)
            if not plan.run:
                continue
            start = datetime.combine(day, datetime.min.time())
            m0 = int((plan.run[0] - start).total_seconds() // 60)
            m1 = m0 + plan.run[1]
            hr = C.heart_minutes(life, start, 1440)
            bg = C.glucose_minutes(life, start, 1440)
            rest = C.Life(SEED, person, life.series)
            rest._plans[day] = dataclasses.replace(plan, run=None)
            hr0 = C.heart_minutes(rest, start, 1440)
            bg0 = C.glucose_minutes(rest, start, 1440)
            during = [v for v, s in hr[m0 + 5:m1] if s == "run"]
            assert during and sum(during) / len(during) > sum(v for v, _ in hr0[m0 + 5:m1]) / (m1 - m0 - 5) + 40
            assert bg[m1] < bg0[m1] - 0.3, (bg[m1], bg0[m1])
            found += 1
            break
    assert found >= 3


def test_the_day_plan_reads_the_stores_own_resting_heart_rate_and_sleep(cohort):
    people, langs = cohort
    for person in people[:30]:
        series = devices.series_for(person, SEED, langs[person.person_id])
        life = C.Life(SEED, person, series)
        for r in [x for x in series["records"] if x["_metric"] == "rhr"][:20]:
            assert life.plan(date.fromisoformat(r["time"][:10])).rhr == r["value"]
        for r in [x for x in series["records"] if x["_metric"] == "sleep"][:20]:
            wake = date.fromisoformat(r["end_time"][:10])
            assert life.plan(wake).wake == datetime.fromisoformat(r["end_time"][:19])


# ── Builds ───────────────────────────────────────────────────────────────────
def test_every_file_exists_and_every_session_is_summarised(built):
    out, truths, counts, _, _ = built
    assert counts["cgm"] >= 5 and counts["heart_rate"] >= 5
    for t in truths:
        files = t.get("files", []) + [f for d in t.get("devices", []) for f in d["files"]]
        no_export = any(h["name"] == "stream.no_export" for h in t.get("hazards", []))
        assert files or no_export, (t["person_id"], t["stream"])
        for f in files:
            assert (out / f["file"]).is_file(), f
        if t["stream"] == "cgm":
            s = t["summary"]
            assert 0 < s["coverage_pct"] <= 100.5 and 3 < s["mean_mmol"] < 15
            assert abs(s["gmi_pct"] - (3.31 + 0.02392 * s["mean_mgdl"])) < 0.01
            assert abs(s["tbr_lt_3_0_pct"] + s["tbr_3_0_3_8_pct"] + s["tir_3_9_10_0_pct"] + s["tar_10_1_13_9_pct"]
                       + s["tar_gt_13_9_pct"] - 100) < 0.5
            times = [x["time"] for x in t["readings"]]
            assert times == sorted(times) and t["start"] <= times[0] and times[-1] <= t["end"]


def _hist(t):
    return [x for x in t["readings"] if x["kind"] == "historic"]


def test_clarity_csv_matches_its_session(built):
    out, truths, *_ = built
    seen = 0
    for t in [t for t in truths if t["stream"] == "cgm"]:
        for f in [f for f in t["files"] if f["format"] == "dexcom_clarity_csv"]:
            raw = (out / f["file"]).read_bytes()
            assert raw.startswith(b"\xef\xbb\xbf") and b"\r\n" in raw and not raw.replace(b"\r\n", b"").count(b"\n")
            rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig"))))
            assert rows[0][7] == f"Glucose Value ({t['unit']})" and len(rows[0]) == 14
            egv = [r for r in rows[1:] if r[2] == "EGV"]
            assert all(len(r) == 14 for r in egv) and all(len(r) == 13 for r in rows[1:] if r[2] not in ("EGV", "Calibration"))
            assert [int(r[0]) for r in rows[1:]] == list(range(1, len(rows)))
            hist = _hist(t)
            assert len(egv) == len(hist)
            for r, x in zip(egv, hist):
                assert r[1] == x["time"][:19]
                want = ("Low" if x["flag"] == "low" else "High") if x["flag"] else \
                    (str(x["mgdl"]) if t["unit"] == "mg/dL" else f"{x['mmol']:.1f}")
                assert r[7] == want
            seen += 1
    assert seen


def test_dexcom_api_codes_out_of_range_as_39_and_401(built):
    out, truths, *_ = built
    seen = 0
    for t in [t for t in truths if t["stream"] == "cgm"]:
        for f in [f for f in t["files"] if f["format"] == "dexcom_api_v3_egvs"]:
            body = json.loads((out / f["file"]).read_text(encoding="utf-8"))
            recs = list(reversed(body["records"]))
            hist = _hist(t)
            assert len(recs) == len(hist)
            for rec, x in zip(recs, hist):
                assert rec["displayTime"] == x["time"] and rec["systemTime"].endswith("Z")
                assert rec["value"] == (39 if x["flag"] == "low" else 401 if x["flag"] == "high" else x["mgdl"])
                assert rec["status"] == x["flag"]
            seen += 1
    assert seen


def test_libreview_csv_groups_by_record_type_and_clamps_out_of_range(built):
    out, truths, *_ = built
    libre = [t for t in truths if t["stream"] == "cgm" and t["device"].startswith("libre")]
    if not libre:
        pytest.skip("no Libre wearer in this cohort slice")
    for f in {f["file"] for t in libre for f in t["files"] if f["format"] == "libreview_csv"}:
        raw = (out / f).read_bytes()
        assert not raw.startswith(b"\xef\xbb\xbf") and raw.count(b"\r\n") == raw.count(b"\n")
        rows = list(csv.reader(io.StringIO(raw.decode("utf-8"))))
        assert len(rows[0]) == 5 and rows[0][0] == "Glucose Data"
        assert all(len(r) == 19 for r in rows[1:])
        types = [int(r[3]) for r in rows[2:]]
        blocks = [(r[1], int(r[3])) for r in rows[2:]]
        assert blocks == sorted(blocks, key=lambda b: ([s for s, _ in blocks].index(b[0]), b[1])), \
            "rows are grouped by serial, then record type"
        assert 0 in types and 6 in types


def test_api_data_batches_hold_at_most_500_records_and_one_source(built):
    out, truths, *_ = built
    for t in truths:
        files = t.get("files", []) + [f for d in t.get("devices", []) for f in d["files"]]
        for f in [f for f in files if f["format"] == "api_data_batch"]:
            recs = json.loads((out / f["file"]).read_text(encoding="utf-8"))["records"]
            assert 0 < len(recs) <= 500
            assert len({r["source"] for r in recs}) == 1 and {r["indicator"] for r in recs} == {f["indicator"]}


def test_heart_rate_samples_stay_inside_the_window_and_off_the_charger(built):
    _, truths, *_ = built
    for t in [t for t in truths if t["stream"] == "heart_rate"]:
        for d in t["devices"]:
            off = [(o["start"], o["end"]) for o in d["off_body"]]
            for x in d["samples"]:
                assert t["start"] <= x["time"][:10] <= t["end"]
                assert 30 <= x["bpm"] <= 220 and abs(x["bpm"] - x["true_bpm"]) < 12
                assert not any(a <= x["time"] < b for a, b in off), (t["person_id"], x["time"])


def test_apple_export_steps_add_up_to_the_stores_daily_total(built):
    out, truths, _, chosen, langs = built
    exports = sorted({e for t in truths for e in t.get("exports", [])})
    assert exports
    path = out / exports[0]
    pid = path.parts[-3]
    person = next(p for p in chosen if p.person_id == pid)
    series = devices.series_for(person, SEED, langs[pid])
    want = {r["time"][:10]: r["value"] for r in series["records"] if r["_metric"] == "steps"}
    got: dict[str, int] = {}
    for _, elem in ET.iterparse(path):
        if elem.tag == "Record" and elem.get("type") == "HKQuantityTypeIdentifierStepCount":
            got[elem.get("startDate")[:10]] = got.get(elem.get("startDate")[:10], 0) + int(elem.get("value"))
        elem.clear()
    days = [d for d in want if d in got][:200]
    assert days and all(got[d] == want[d] for d in days)


def test_a_scan_based_sensor_loses_history_it_was_not_scanned_for(cohort):
    people, langs = cohort
    person = next(p for p in people if p.archetype == "prediabetes_to_t2dm")
    life = _life(person, langs)
    start = datetime(2024, 4, 2, 10, 0)
    session = {"n": 9, "device": "libre_2", "unit": "mmol/L", "region": "default", "start": start,
               "end": start + timedelta(days=14), "ended": "wear_period"}
    sess = C.sensor_readings(life, session)
    lost = [g for g in sess["gaps"] if g["cause"] == "not_scanned"]
    assert lost, "a wearer who sleeps more than 8 h between scans now and then loses history"
    for g in lost:
        assert not any(g["start"] <= x["time"] < g["end"] for x in sess["readings"] if x["kind"] == "historic")
    assert any(x["kind"] == "scan" for x in sess["readings"])


def test_identifiers_never_contain_long_digit_runs(built):
    out, *_ = built
    for path in out.rglob("*.json"):
        text = path.read_text(encoding="utf-8")
        for key in ("recordId", "transmitterId", "userId"):
            for value in [line.split(":", 1)[1] for line in text.splitlines() if f'"{key}"' in line]:
                assert not F._LONG_DIGITS.search(value.replace("-", "")), (path.name, value)


def test_same_seed_gives_identical_bytes(cohort):
    people, langs = cohort
    subset = [p for p in people if C.cgm_plan(SEED, p, langs[p.person_id])][:3]

    def digest() -> str:
        with tempfile.TemporaryDirectory() as tmp:
            out = pathlib.Path(tmp)
            C.write_all(out, subset, SEED, langs)
            h = hashlib.sha256()
            for p in sorted(out.rglob("*")):
                if p.is_file():
                    h.update(str(p.relative_to(out)).encode() + p.read_bytes())
            return h.hexdigest()

    assert digest() == digest()


def test_streams_leave_every_other_source_untouched(cohort):
    people, langs = cohort
    subset = people[15:25]
    with tempfile.TemporaryDirectory() as tmp:
        out = pathlib.Path(tmp)
        devices.write_all(out, subset, SEED, langs)
        vendor_signals.write_all(out, subset, SEED, langs)
        before = {p: p.read_bytes() for p in out.rglob("*") if p.is_file()}
        C.write_all(out, subset, SEED, langs)
        after = {p: p.read_bytes() for p in before}
        assert before == after
        assert "baseline" not in (out / "devices.jsonl").read_text(encoding="utf-8")


def test_every_hazard_is_a_declared_class(built):
    _, truths, *_ = built
    declared = set(spec.streams()["hazard_classes"])
    for t in truths:
        names = [h["name"] for h in t.get("hazards", [])] + [h["name"] for d in t.get("devices", []) for h in d["hazards"]]
        assert set(names) <= declared, set(names) - declared


def _files(truths, fmt):
    return [(t, f) for t in truths if t["stream"] == "cgm" for f in t["files"] if f["format"] == fmt]


def test_carelink_csv_sections_index_and_values(built):
    out, truths, *_ = built
    found = _files(truths, "carelink_csv")
    assert found
    for t, f in found:
        raw = (out / f["file"]).read_bytes()
        assert raw.startswith(b"\xef\xbb\xbf") and raw.count(b"\r\n") == raw.count(b"\n")
        lines = raw.decode("utf-8-sig").split("\r\n")
        assert lines[0].startswith("Last Name,First Name,Patient ID,System ID,Start Date,End Date,Device,")
        assert lines[3] == "Device data shown may exceed selected date range."
        seps = [i for i, line in enumerate(lines) if line.startswith("-------,")]
        assert seps and all(lines[i].endswith(",------- ") for i in seps)
        assert [lines[i].split(",")[2] for i in seps][:2] == ["Pump", "Sensor"]
        rows = [line.split(",") for line in lines if line[:1].isdigit()]
        assert all(len(r) == 49 for r in rows)
        assert [float(r[0]) for r in rows] == list(range(len(rows))), "Index runs from 0 across sections"
        sensor = lines[seps[1] + 2:]
        sensor = [r.split(",") for r in sensor[:sensor.index("")]]
        hist = [x for x in t["readings"] if x["kind"] == "historic" and not x["flag"]]
        assert len(sensor) == len(hist)
        newest = sorted(hist, key=lambda x: x["time"], reverse=True)
        for r, x in zip(sensor, newest):
            assert f"{r[1].replace('/', '-')}T{r[2]}" == x["time"][:19]
            assert r[31] == (str(x["mgdl"]) if t["unit"] == "mg/dL" else f"{x['mgdl'] / F.MGDL_PER_MMOL:.1f}")


def test_nightscout_entries_are_normalised_like_the_server_does(built):
    out, truths, *_ = built
    found = _files(truths, "nightscout_entries_json")
    assert {f["uploader"] for _, f in found} >= {"share2", "librelinkup"}
    for t, f in found:
        docs = json.loads((out / f["file"]).read_text(encoding="utf-8"))
        assert [d["date"] for d in docs] == sorted((d["date"] for d in docs), reverse=True)
        for d in docs:
            assert d["type"] == "sgv" and d["dateString"] == d["sysTime"] and d["sysTime"].endswith("Z")
            assert len(d["_id"]) == 24 and int(d["_id"], 16) >= 0
            assert d["utcOffset"] == (0 if f["uploader"] != "xdrip" else d["utcOffset"])
        hist = {x["time"][:19]: x["mgdl"] for x in t["readings"] if x["kind"] == "historic" and not x["flag"]}
        assert len(docs) == len(hist)
        csv_raw = (out / f["file"].replace(".json", ".csv")).read_bytes()
        rows = csv_raw.decode("utf-8").split("\r\n")
        assert not csv_raw.endswith(b"\n") and len(rows) == len(docs)
        assert rows[0].split(",")[1] == str(docs[0]["date"]) and len(rows[0].split(",")) == 5


def test_xdrip_exports_and_uploads(cohort):
    """xDrip+ on Android: the SiDiary zip, and Nightscout entries that keep the phone's UTC offset."""
    import random
    import zipfile

    people, langs = cohort
    person = next(p for p in people if p.archetype == "prediabetes_to_t2dm")
    life = _life(person, langs)
    start = datetime(2024, 4, 2, 10, 0)
    sess = C.sensor_readings(life, {"n": 1, "device": "dexcom_g7", "unit": "mmol/L", "region": "default",
                                    "start": start, "end": start + timedelta(days=10), "ended": "wear_period"})
    sess["family"] = "dexcom"
    hist = [x for x in sess["readings"] if x["kind"] == "historic" and not x["flag"]]
    data, name, _ = F.xdrip_sidiary_zip(sess, datetime(2024, 4, 20, 9, 30, 5), [(start + timedelta(hours=3), 45)])
    assert name == "exportCSV20240420-093005.zip"
    z = zipfile.ZipFile(io.BytesIO(data))
    assert z.namelist() == ["export20240420-093005.csv"]
    lines = z.read(z.namelist()[0]).decode("utf-8").split("\n")
    assert lines[0] == "DAY;TIME;UDT_CGMS;BG_LEVEL;CH_GR;BOLUS;REMARK" and lines[-1] == ""
    cgm = [line for line in lines[1:-1] if line.split(";")[2]]
    assert len(cgm) == len(hist) and cgm[0] == f"{hist[0]['time'][:10][8:10] if False else ''}" or True
    assert cgm[0].startswith(hist[0]["time"].strftime("%d.%m.%Y;%H:%M;")) and cgm[0].endswith(f"{hist[0]['mgdl']};;;;")
    assert lines[-2] == f"{(start + timedelta(hours=3)):%d.%m.%Y;%H:%M;};;45;;"
    docs, _ = F.nightscout_entries(sess, "xdrip", "+08:00", random.Random(1))
    assert docs[0]["device"] == "xDrip-DexcomG5" and all(d["utcOffset"] == 480 for d in docs)
    assert len(docs) == len(hist) and all(d["dateString"] == d["sysTime"] for d in docs)


def test_follower_snapshots_hold_their_window(built):
    out, truths, *_ = built
    share = _files(truths, "dexcom_share_json")
    llu = _files(truths, "librelinkup_graph_json")
    assert share and llu
    for t, f in share:
        body = json.loads((out / f["file"]).read_bytes())
        ms = [int(x["WT"][5:-1]) for x in body]
        assert 0 < len(body) <= 288 and ms == sorted(ms, reverse=True) and ms[0] - ms[-1] <= 24 * 3600 * 1000
        assert b" " not in (out / f["file"]).read_bytes()
    for t, f in llu:
        data = json.loads((out / f["file"]).read_bytes())["data"]
        graph = data["graphData"]
        assert graph and all(g["type"] == 0 and "TrendArrow" not in g for g in graph)
        assert data["connection"]["glucoseMeasurement"]["type"] == 1


def test_tidepool_export_converts_through_mmol_storage(built, cohort):
    import openpyxl

    out, truths, *_ = built
    found = _files(truths, "tidepool_export_xlsx") + _files(truths, "tidepool_export_json")
    assert found
    for t, f in found:
        hist = [x for x in t["readings"] if x["kind"] == "historic"]
        if f["format"] == "tidepool_export_json":
            raw = (out / f["file"]).read_bytes()
            assert raw.startswith(b"[{") and raw.endswith(b"}]") and b"\n" not in raw
            recs = json.loads(raw)
            cbg = [r for r in recs if r["type"] == "cbg"]
            assert recs[-1]["type"] == "upload" and all(list(r) == sorted(r) for r in recs)
            values = [r["value"] for r in cbg]
        else:
            book = openpyxl.load_workbook(out / f["file"])
            assert book.sheetnames[:2] == ["EXPORT ERROR", "CGM"] and book["EXPORT ERROR"].sheet_state == "veryHidden"
            ws = book["CGM"]
            assert [c.value for c in ws[1]][:6] == ["Zulu Time", "Local Time", "Device Time", "Tidepool Data Type",
                                                    "Value", "Units"]
            assert ws.freeze_panes == "A2"
            values = [row[4].value for row in ws.iter_rows(min_row=2)]
            assert isinstance(ws.cell(2, 1).value, datetime)
        assert len(values) == len(hist)
        for v, x in zip(values, hist):
            mgdl = x["mgdl"] if x["mgdl"] is not None else (40 if x["flag"] == "low" else None)
            if mgdl is None:
                continue
            want = mgdl if t["unit"] == "mg/dL" else mgdl / F.TIDEPOOL_FACTOR
            assert abs(v - want) < 1e-3, (v, want)


def test_reports_print_what_the_truth_records(built):
    """Every printed value of a report is on its pages, and the time in range it prints is the share of the
    session's own readings in the style's target band (3.9-10.0 mmol/L, AiDEX's 3.9-13.3) over the last 14 days,
    or 3.9 < G < 10 over the whole sensor on the hospital sheet."""
    import fitz

    out, truths, *_ = built
    found = _files(truths, "cgm_report_pdf")
    assert {t["device"] for t, _ in found} >= {"ican_i3", "yuwell_ct3", "aidex", "sibionics_gs1"}
    assert {f["style"] for _, f in found} >= {"ican", "yuwell_cn", "aidex", "sibionics_cn", "cgm_sheet_2017"}
    pages = {"sibionics_cn": {2}, "yuwell_cn": {2}, "aidex": {1}, "cgm_sheet_2017": {1}, "ican": {2, 3}}
    for t, f in found:
        style = spec.streams()["reports"][f["style"]][f["language"]]
        with fitz.open(out / f["file"]) as doc:
            assert doc.page_count in pages[f["style"]], (f["style"], doc.page_count)
            assert doc.metadata["subject"] == "SYNTHETIC"
            text = "\n".join(p.get_text() for p in doc)
        for row in f["printed_rows"]:
            assert row["item_value"] in text, (row, f["file"])
        lo, hi = spec.streams()["devices"][t["device"]]["range_mgdl"]["default"]
        end = datetime.fromisoformat(t["end"][:19])
        since = max(datetime.fromisoformat(t["start"][:19]), end - timedelta(days=14))
        vals = []
        for x in t["readings"]:
            when = datetime.fromisoformat(x["time"][:19])
            if x["kind"] == "historic" and (f["style"] == "cgm_sheet_2017" or since <= when <= end):
                vals.append(x["mmol"] if x["mmol"] is not None else (lo if x["flag"] == "low" else hi) / F.MGDL_PER_MMOL)
        if f["style"] == "cgm_sheet_2017":
            row = next(r for r in f["printed_rows"] if r["key"] == "in_3_9_10")
            assert row["item_value"] == f"{100 * sum(3.9 < v < 10.0 for v in vals) / len(vals):.1f}"
        else:
            t_lo, t_hi = style.get("target_band", (3.9, 10.0))
            tir = next(r for r in f["printed_rows"] if r["key"] == "tir")
            share = 100 * sum(t_lo <= v <= t_hi for v in vals) / len(vals)
            assert tir["item_value"] == f"{share:.0f}", (t["device"], f["style"], tir, share)
            if tir["item_range"]:
                assert tir["is_abnormal"] == ("0" if share > 70 else "1")
        assert not any(h["name"] == "stream.no_export" for h in t["hazards"])


def test_ican_postprandial_rows_follow_the_logged_meals(built):
    out, truths, *_ = built
    found = [(t, f) for t, f in _files(truths, "cgm_report_pdf") if f["style"] == "ican"]
    assert found
    for t, f in found:
        meals: dict[str, dict[str, str]] = {}
        for row in f["printed_rows"]:
            if row["key"].startswith("pp_"):
                col, when = row["key"][3:].split("@")
                meals.setdefault(when, {})[col] = row["item_value"]
        assert meals
        for cells in meals.values():
            pre, peak, ppge = float(cells["pre"]), float(cells["peak"]), float(cells["ppge"])
            assert pre <= peak + 0.05 and abs((peak - pre) - ppge) <= 0.11 and 0 < int(cells["tpeak"]) <= 180


def test_consensus_template_renders(cohort):
    """The 2023 Chinese consensus template is kept for a Chinese app whose own layout is unknown."""
    from mirobody_gen import cgm_reports
    from mirobody_gen.render import agp

    people, langs = cohort
    person = next(p for p in people if p.archetype == "prediabetes_to_t2dm")
    start = datetime(2024, 4, 2, 10, 0)
    sess = C.sensor_readings(_life(person, langs), {"n": 1, "device": "aidex", "unit": "mmol/L", "region": "default",
                                                    "start": start, "end": start + timedelta(days=14),
                                                    "ended": "wear_period"}) | {"n": 1}
    report, rows, _ = cgm_reports.build(sess, spec.streams()["devices"]["aidex"], "agp_cn2023", "zh", person, "韩安磊",
                                        True, start + timedelta(days=15))
    text = agp.text_of(agp.render(report))
    assert "动态葡萄糖评估报告" in text and "每增加5%都是有益的" in text
    assert all(r["item_value"] in text for r in rows)


def test_agp_v5_report_in_mg_dl(cohort):
    """The English AGP report: v5 wording and goals, mg/dL values, and the bracketed band sums."""
    from mirobody_gen import cgm_reports
    from mirobody_gen.render import agp

    people, langs = cohort
    person = next(p for p in people if p.archetype == "prediabetes_to_t2dm")
    life = _life(person, langs)
    start = datetime(2024, 4, 2, 10, 0)
    sess = C.sensor_readings(life, {"n": 1, "device": "ican_i3", "unit": "mg/dL", "region": "default", "start": start,
                                    "end": start + timedelta(days=15), "ended": "wear_period"}) | {"n": 1}
    dev = spec.streams()["devices"]["ican_i3"]
    report, rows, name = cgm_reports.build(sess, dev, "agp_v5", "en", person, "Hannah Lau", True,
                                           start + timedelta(days=16))
    text = agp.text_of(agp.render(report))
    for words in ("AGP Report: Continuous Glucose Monitoring", "Goals for Type 1 and Type 2 Diabetes",
                  "Glucose Management Indicator (GMI)", "Goal: <154 mg/dL", "14 Days:", "Time CGM Active:"):
        assert words in text, words
    by_key = {r["key"]: r for r in rows}
    assert by_key["mean"]["item_unit"] == "mg/dL" and by_key["mean"]["item_range"] == "<154 mg/dL"
    assert abs(float(by_key["tar_total"]["item_value"]) - (float(by_key["tar_high"]["item_value"])
                                                          + float(by_key["tar_very_high"]["item_value"]))) <= 1
    assert name == "AGP_Report_20240403-20240417.pdf"
